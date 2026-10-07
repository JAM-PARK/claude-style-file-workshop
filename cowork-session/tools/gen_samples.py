"""samples/ 아래 실습 샘플 네 묶음과 cowork-samples.zip을 만든다.

실행: tools/.venv/bin/python tools/gen_samples.py
원본 데이터는 tools/sample_data.py. 정답지는 tools/gen_answer_keys.py가 같은 데이터로 만든다.
폰트는 tools/fonts-src/ (tools/fetch-fonts.sh)를 쓴다.
"""
import datetime as dt
import io
import os
import pathlib
import random
import shutil
import zipfile

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

import sample_data as D

ROOT = pathlib.Path(__file__).resolve().parent.parent
FONTS = ROOT / 'tools/fonts-src'
OUT = ROOT / 'samples'
DIRS = {
    'r': OUT / '01_영수증',
    'd': OUT / '02_다운로드폴더_정리전',
    'q': OUT / '03_견적서비교',
    'm': OUT / '04_회의녹취',
}
rng = random.Random(20260907)
# 다시 만들어도 파일이 바이트 단위로 같도록 문서 안 생성 시각을 고정한다
FIXED = dt.datetime(2026, 9, 1, 9, 0, 0)


def won(n):
    return f'{n:,}'


def font(size, bold=False, mono=False):
    name = ('NanumGothicCoding-Bold' if bold else 'NanumGothicCoding-Regular') if mono else \
        ('NanumGothic-Bold' if bold else 'NanumGothic-Regular')
    return ImageFont.truetype(str(FONTS / f'{name}.ttf'), size)


def set_mtime(path, date, time='10:00'):
    t = dt.datetime.strptime(f'{date} {time}', '%Y-%m-%d %H:%M').timestamp()
    os.utime(path, (t, t))


# ---------------------------------------------------------------------------
# 01 영수증
# ---------------------------------------------------------------------------
W_RECEIPT = 860


def receipt_lines(r):
    L = []
    if r['kind'] == 'taxi':
        L += [('c', '택 시 영 수 증', True), ('-',), ('l', f"차량번호  {r['car']}"),
              ('l', f"사업자번호 {r['biz']}"), ('-',),
              ('l', f"승차  {r['date']} {r['time']}"), ('l', f"하차  {r['date']} {r['end']}"), ('-',),
              ('lr', '요금', f"{won(r['amount'])}원", True), ('-',),
              ('l', f"결제  {'카드 (끝자리 0000)' if r['pay'] == '카드' else '현금'}"),
              ('l', '승인번호 00000000' if r['pay'] == '카드' else '현금 결제'), ('-',),
              ('c', '이용해 주셔서 감사합니다')]
        return L
    total = D.receipt_total(r)
    vat = round(total / 11)
    L += [('c', r['store'], True), ('c', f"사업자번호 {r['biz']}"), ('c', f"{D.FAKE_ADDR} {int(r['biz'][-3:]) % 90 + 10}"),
          ('c', f"전화 {r['tel']}"), ('-',), ('l', f"{r['date']} {r['time']}"), ('-',),
          ('lr3', '품목', '수량', '금액')]
    for name, q, p in r['items']:
        L.append(('lr3', name, str(q), won(q * p)))
    L += [('-',), ('lr', '합계', f"{won(total)}원", True), ('lr', '  공급가액', won(total - vat)), ('lr', '  부가세', won(vat)),
          ('-',)]
    if r['pay'] == '카드':
        L += [('l', '결제  신용카드 (끝자리 0000)'), ('l', '승인번호 00000000'), ('l', '할부  일시불')]
    else:
        L += [('l', '결제  현금'), ('l', f"받은 돈 {won(total)}원")]
    L += [('-',), ('c', '감사합니다')]
    return L


def draw_receipt(r):
    f, fb = font(36, mono=True), font(39, bold=True, mono=True)
    lines = receipt_lines(r)
    lh = 56
    h = 60 + lh * len(lines) + 40
    im = Image.new('RGB', (W_RECEIPT, h), (250, 249, 245))
    d = ImageDraw.Draw(im)
    y = 40
    for ln in lines:
        kind = ln[0]
        if kind == '-':
            d.text((40, y), '-' * 40, font=f, fill=(90, 90, 90))
        elif kind == 'c':
            ff = fb if len(ln) > 2 and ln[2] else f
            w = d.textlength(ln[1], font=ff)
            d.text(((W_RECEIPT - w) / 2, y), ln[1], font=ff, fill=(25, 25, 25))
        elif kind == 'l':
            d.text((40, y), ln[1], font=f, fill=(25, 25, 25))
        elif kind == 'lr':
            ff = fb if len(ln) > 3 and ln[3] else f
            d.text((40, y), ln[1], font=ff, fill=(25, 25, 25))
            w = d.textlength(ln[2], font=ff)
            d.text((W_RECEIPT - 40 - w, y), ln[2], font=ff, fill=(25, 25, 25))
        elif kind == 'lr3':
            d.text((40, y), ln[1], font=f, fill=(25, 25, 25))
            d.text((560, y), ln[2], font=f, fill=(25, 25, 25))
            w = d.textlength(ln[3], font=f)
            d.text((W_RECEIPT - 40 - w, y), ln[3], font=f, fill=(25, 25, 25))
        y += lh
    return im


def photo(receipt_img, tilt, blur=0.0, seed=0):
    """책상 위에 놓고 찍은 사진처럼."""
    r = random.Random(seed)
    bg = Image.new('RGB', (1200, 1600), (r.randint(150, 175), r.randint(140, 160), r.randint(120, 140)))
    noise = Image.effect_noise((1200, 1600), 18).convert('RGB')
    bg = Image.blend(bg, noise, 0.08)
    rot = receipt_img.convert('RGBA').rotate(tilt, expand=True, resample=Image.BICUBIC)
    x = (1200 - rot.width) // 2 + r.randint(-40, 40)
    y = max(20, (1600 - rot.height) // 2 + r.randint(-30, 30))
    bg.paste(rot, (x, y), rot)
    out = bg.filter(ImageFilter.GaussianBlur(0.5 + blur))
    return out


def draw_online(r):
    """휴대폰 결제 완료 화면 캡처."""
    W, H = 720, 1400
    im = Image.new('RGB', (W, H), (255, 255, 255))
    d = ImageDraw.Draw(im)
    f, fb, fs = font(30), font(34, bold=True), font(24)
    d.rectangle([0, 0, W, 60], fill=(240, 240, 240))
    d.text((24, 16), f"{r['time']}", font=fs, fill=(40, 40, 40))
    d.text((24, 100), '결제가 완료되었습니다', font=font(40, bold=True), fill=(20, 20, 20))
    d.text((24, 160), f"{r['store']} 주문", font=f, fill=(90, 90, 90))
    y = 250
    rows = [('주문번호', r['order']), ('결제일시', f"{r['date']} {r['time']}"), ('판매자', r['store']),
            ('사업자번호', r['biz'])]
    for k, v in rows:
        d.text((24, y), k, font=f, fill=(110, 110, 110))
        d.text((260, y), v, font=f, fill=(20, 20, 20))
        y += 56
    d.line([24, y + 10, W - 24, y + 10], fill=(220, 220, 220), width=2)
    y += 40
    for name, q, p in r['items']:
        d.text((24, y), name, font=f, fill=(20, 20, 20))
        s = f'{won(q * p)}원'
        d.text((W - 24 - d.textlength(s, font=f), y), s, font=f, fill=(20, 20, 20))
        y += 56
    d.line([24, y + 10, W - 24, y + 10], fill=(220, 220, 220), width=2)
    y += 40
    s = f'{won(D.receipt_total(r))}원'
    d.text((24, y), '총 결제금액', font=fb, fill=(20, 20, 20))
    d.text((W - 24 - d.textlength(s, font=fb), y), s, font=fb, fill=(20, 20, 20))
    y += 70
    d.text((24, y), '결제수단  신용카드 (끝자리 0000)', font=f, fill=(90, 90, 90))
    d.rectangle([24, H - 140, W - 24, H - 60], outline=(200, 200, 200), width=2)
    t = '주문 상세 보기'
    d.text(((W - d.textlength(t, font=f)) / 2, H - 118), t, font=f, fill=(60, 60, 60))
    return im


def build_receipts():
    out = DIRS['r']
    by_id = {r['id']: r for r in D.RECEIPTS}
    for i, r in enumerate(D.RECEIPTS):
        src = by_id[r['dup_of']] if 'dup_of' in r else r
        if src['kind'] == 'online':
            draw_online(src).save(out / r['file'])
        else:
            img = photo(draw_receipt(src), r.get('tilt', 0), r.get('blur', 0), seed=i)
            img.save(out / r['file'], quality=86)
        date = r['file'][4:12] if r['file'].startswith('IMG') else r['file'][11:19]
        set_mtime(out / r['file'], f'{date[:4]}-{date[4:6]}-{date[6:]}', '20:00')


# ---------------------------------------------------------------------------
# PDF 공통
# ---------------------------------------------------------------------------
pdfmetrics.registerFont(TTFont('NG', str(FONTS / 'NanumGothic-Regular.ttf')))
pdfmetrics.registerFont(TTFont('NGB', str(FONTS / 'NanumGothic-Bold.ttf')))


def pdf_doc(path, title, blocks, footnotes=(), small_footnotes=True):
    """blocks: ('p', text) / ('h', text) / ('table', rows, col_widths_mm) / ('gap', mm)"""
    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    W, H = A4
    x0, y = 20 * mm, H - 25 * mm
    c.setFont('NGB', 20)
    c.drawString(x0, y, title)
    y -= 12 * mm
    for b in blocks:
        if b[0] == 'gap':
            y -= b[1] * mm
        elif b[0] in ('p', 'h'):
            c.setFont('NGB' if b[0] == 'h' else 'NG', 12 if b[0] == 'h' else 10.5)
            for line in b[1].split('\n'):
                c.drawString(x0, y, line)
                y -= 6.5 * mm
        elif b[0] == 'table':
            rows, widths = b[1], [w * mm for w in b[2]]
            for ri, row in enumerate(rows):
                c.setFont('NGB' if ri == 0 else 'NG', 10)
                x = x0
                for ci, cell in enumerate(row):
                    if ci > 0 and isinstance(cell, str) and cell.replace(',', '').replace('원', '').isdigit():
                        c.drawRightString(x + widths[ci] - 2 * mm, y, cell)
                    else:
                        c.drawString(x + 2 * mm, y, str(cell))
                    x += widths[ci]
                c.setLineWidth(0.4)
                c.line(x0, y - 2.5 * mm, x0 + sum(widths), y - 2.5 * mm)
                y -= 8 * mm
            y -= 2 * mm
    if footnotes:
        fy = 30 * mm + 4.2 * mm * len(footnotes)
        c.setFont('NG', 7.5 if small_footnotes else 10)
        for fn in footnotes:
            c.drawString(x0, fy, '※ ' + fn)
            fy -= 4.2 * mm
    c.showPage()
    c.save()


# ---------------------------------------------------------------------------
# 03 견적서비교
# ---------------------------------------------------------------------------
def build_quotes():
    out = DIRS['q']
    (out / '요청사항.txt').write_text(
        f'{D.SHOP} 홈페이지 제작 요청사항 (업체 세 곳에 같은 내용으로 보냄)\n\n'
        f"- 페이지: {D.REQUEST['pages']}\n- 기능: {D.REQUEST['feature']}\n- 비교 기준: {D.REQUEST['period']}\n"
        '- 휴대폰에서도 잘 보여야 함\n- 사진은 우리가 찍은 것을 씀\n', encoding='utf-8')
    set_mtime(out / '요청사항.txt', '2026-09-02')
    for q in D.QUOTES:
        rows = [('항목', '금액(원)')] + [(n, won(a)) for n, a in q['lines']]
        blocks = [
            ('p', f"수신: {D.SHOP} 귀하\n견적일: {q['date']}    유효기간: {q['valid']}까지"),
            ('gap', 2),
            ('p', f"공급자: {q['vendor']}   사업자번호 {q['biz']}\n주소: {D.FAKE_ADDR} {200 + ord(q['id'])}   전화 {q['tel']}"),
            ('gap', 4), ('h', '견적 내용'), ('table', rows, (130, 40)),
        ]
        if q.get('options'):
            blocks += [('h', '선택 옵션'), ('table', [('옵션', '금액(원)')] + [(n, won(a)) for n, a in q['options']], (130, 40))]
        blocks += [('gap', 2), ('h', f"{q['headline_note']}: {won(D.quote_headline(q))}원")]
        pdf_doc(out / q['file'], '견 적 서', blocks, q['footnotes'])
        set_mtime(out / q['file'], q['date'], '17:00')


# ---------------------------------------------------------------------------
# 04 회의녹취
# ---------------------------------------------------------------------------
TRANSCRIPT = [
    ('윤하늘', '자, 다들 왔죠? 음, 오늘은 추석 선물세트랑 홈페이지 얘기 두 개만 하고 끝낼게요. 세린 님 바쁘신데 와주셔서 감사해요.'),
    ('오세린', '아니에요, 저도 라벨 방향 한 번 맞춰야 해서요.'),
    ('박소이', '저 커피 하나만 타 올게요. 아, 아니다, 그냥 할게요.'),
    ('윤하늘', '그래요. 어, 일단 선물세트부터. 작년에는 세 종류 했잖아요. 근데 솔직히 세 개는 너무 힘들었어.'),
    ('박소이', '네, 작년에 파운드 세트 포장하다가 밤 열한 시까지 있었어요.'),
    ('윤하늘', '그쵸. 그래서 올해는 두 개로 줄이려고. 쿠키 세트 하나, 파운드 세트 하나.'),
    ('오세린', '쿠키 세트는 몇 개 들어가요?'),
    ('윤하늘', '열두 개요. 버터쿠키, 얼그레이, 초코칩 네 개씩. 가격은 이만 팔천 원 생각하고 있어요.'),
    ('박소이', '이만 팔천이요? 작년엔 이만 육천이었는데.'),
    ('윤하늘', '버터가 너무 올랐어. 음, 이천 원은 올려야 돼요. 대신 상자를 좀 예쁘게.'),
    ('오세린', '네네, 그거 좋아요. 가격 올릴 때 포장이 바뀌면 덜 아깝게 느끼시거든요.'),
    ('윤하늘', '그쵸. 그리고 파운드 세트는 파운드 한 개에 쿠키 여섯 개. 사만 이천 원.'),
    ('박소이', '파운드는 무화과로 해요, 아니면 레몬?'),
    ('윤하늘', '무화과요. 레몬은 추석 느낌이 아니잖아.'),
    ('박소이', '아 맞다, 그렇네요.'),
    ('윤하늘', '그럼 두 종류 이렇게 확정. 쿠키 세트 이만 팔천, 파운드 세트 사만 이천.'),
    ('오세린', '네, 적어 둘게요.'),
    ('윤하늘', '다음은 일정. 추석이 이십오일 금요일이잖아요. 그 전에 받아 가셔야 하니까, 수령은 이십이일 화요일부터 이십사일 목요일까지.'),
    ('박소이', '예약 마감은요?'),
    ('윤하늘', '음, 십구일 토요일. 마감하고 재료 주문 넣어야 돼서.'),
    ('박소이', '십구일 토요일 마감, 이십이일부터 이십사일 수령. 네.'),
    ('오세린', '택배는 안 하세요? 멀리 계신 분들 물어보실 것 같은데.'),
    ('윤하늘', '아, 그게 고민이에요. 작년에 택배 했다가 쿠키 두 박스 깨져서 왔잖아.'),
    ('박소이', '네, 그때 다시 보내드렸어요.'),
    ('윤하늘', '그래서 이번에는, 음, 안 하는 쪽으로 생각하는데. 아직 확정은 아니고. 이건 좀 더 생각해 볼게요.'),
    ('오세린', '네, 그럼 일단 보류로 적을게요.'),
    ('윤하늘', '네. 그 다음, 상자. 소이 씨, 상자 몇 개 남았어요?'),
    ('박소이', '크라프트 상자 중간 사이즈가 한 오십 개? 지난달에 포장마켓에서 오십 개 산 거 거의 그대로 있어요.'),
    ('윤하늘', '그걸로는 모자라지. 작년에 백육십 세트 나갔으니까.'),
    ('박소이', '그럼 이백 개 정도 더 시킬까요?'),
    ('윤하늘', '네, 이백 개. 근데 상자나라 견적이 두 번 왔었죠? 삼백 개짜리랑 이백 개짜리.'),
    ('박소이', '네, 처음에 삼백 개로 받았다가 이백 개로 다시 받았어요. 이백 개짜리로 할게요.'),
    ('윤하늘', '네, 이백 개짜리로. 언제까지 시켜야 돼요?'),
    ('박소이', '제작이 열흘 걸린대요. 그럼 십일까지는 넣어야 할 것 같아요.'),
    ('윤하늘', '그럼 십일 말고 넉넉하게 십일 전에. 소이 씨가 십일까지 발주 부탁해요.'),
    ('박소이', '네, 십일까지 상자 이백 개 발주.'),
    ('윤하늘', '자, 이제 세린 님. 라벨.'),
    ('오세린', '네. 지난번에 보내드린 시안 A, B 보셨죠? 하나는 손글씨 느낌이고 하나는 좀 단정한 거.'),
    ('윤하늘', '봤어요. 저는 B가 좋은데, 소이 씨는 A 좋다 그랬지?'),
    ('박소이', '저는 A요. 손님들이 손글씨 좋아하세요. 근데 B도 깔끔해서 괜찮아요.'),
    ('오세린', '그럼 이렇게 할게요. B 기본으로 하고, A의 손글씨를 이름 부분에만 살짝 넣은 버전. 이렇게 두 개 다시 드릴게요.'),
    ('윤하늘', '오, 좋아요. 언제쯤 될까요?'),
    ('오세린', '음, 다음 주 토요일, 십이일까지 최종 두 개 보내드릴게요.'),
    ('윤하늘', '네, 십이일. 그거 보고 바로 인쇄 넣을게요.'),
    ('오세린', '그리고 리본 색은요? 지금 크라프트 끈이잖아요.'),
    ('박소이', '빨간 리본 하면 추석 느낌 나지 않아요?'),
    ('윤하늘', '음, 빨강은 좀 촌스러울 수도. 근데 라벨 색 보고 정하는 게 맞을 것 같아요.'),
    ('오세린', '네, 라벨 나오면 그 옆에 놓고 보시는 게 좋아요.'),
    ('윤하늘', '그럼 리본은 라벨 보고 다시. 보류.'),
    ('박소이', '리본 보류.'),
    ('윤하늘', '어, 그리고 예약은 어떻게 받죠. 작년에는 전화랑 메모지였잖아. 그거 진짜 헷갈렸어.'),
    ('박소이', '네, 메모지 하나 잃어버려서 손님 한 분 못 챙겨드렸어요.'),
    ('윤하늘', '그러니까. 올해는 표로 받자. 스프레드시트 하나 만들어서 이름, 연락처, 세트 종류, 개수, 수령일.'),
    ('박소이', '제가 만들게요. 언제까지 할까요?'),
    ('윤하늘', '공지 올리기 전에는 있어야 하니까 구일까지.'),
    ('박소이', '네, 구일 수요일까지 예약 표 만들어 둘게요.'),
    ('윤하늘', '그리고 공지는 제가 SNS에 올릴게요. 십일일까지.'),
    ('오세린', '공지에 들어갈 사진은 시안 B 상자 사진 쓰셔도 돼요. 라벨은 최종 아니라고 써 주시고요.'),
    ('윤하늘', '아, 좋아요. 그렇게 할게요.'),
    ('윤하늘', '자, 선물세트는 이 정도면 된 것 같고. 홈페이지.'),
    ('박소이', '아 그거, 견적 다 왔어요?'),
    ('윤하늘', '세 군데서 다 왔어요. 근데 이게 비교가 잘 안 돼. 한 군데는 부가세 포함이고 한 군데는 별도고, 각주에 뭐가 막 써 있고.'),
    ('오세린', '맞아요, 그거 각주를 꼭 보셔야 돼요. 유지보수 비용이 거기 숨어 있는 경우가 많아요.'),
    ('윤하늘', '그래서 그거 제가 이번 주말에 한 번 표로 정리해 보려고요. 다음 주 화요일, 십오일까지는 업체 정할게요.'),
    ('박소이', '홈페이지에서 선물세트 예약도 받을 수 있으면 좋겠는데요.'),
    ('윤하늘', '그게 이번 추석은 못 맞출 것 같아. 홈페이지는 빨라야 시월 오픈이니까. 이번엔 SNS 메시지랑 전화로 받고, 표에 정리하고.'),
    ('박소이', '네, 알겠습니다.'),
    ('오세린', '홈페이지 디자인도 혹시 필요하시면 말씀 주세요. 라벨이랑 톤 맞추면 좋으니까요.'),
    ('윤하늘', '아, 그건 업체 정하고 나서 한번 여쭤볼게요.'),
    ('윤하늘', '음, 그럼 정리하면. 선물세트 두 종류, 이만 팔천이랑 사만 이천. 예약 마감 십구일, 수령 이십이일부터 이십사일.'),
    ('박소이', '상자 이백 개 십일까지 발주, 예약 표 구일까지.'),
    ('오세린', '라벨 최종 두 개 십이일까지.'),
    ('윤하늘', '저는 SNS 공지 십일일, 홈페이지 업체 십오일까지. 리본 색이랑 택배는 보류.'),
    ('박소이', '네.'),
    ('윤하늘', '아, 그리고 세린 님, 오늘 얘기한 거 메일로 한 번 정리해서 보내드릴게요.'),
    ('오세린', '네, 감사합니다. 그럼 저는 이만.'),
    ('윤하늘', '네, 수고하셨어요. 소이 씨는 잠깐 남아서 오늘 반죽 얘기 좀 해요.'),
]


def build_meeting():
    out = DIRS['m']
    t = 12
    lines = [f"[녹취] {D.MEETING['date']} {D.MEETING['place']}",
             '참석: ' + ', '.join(f'{n}({r})' for n, r in D.MEETING['people']),
             '※ 자동 받아쓰기 결과를 그대로 저장한 파일입니다. 말버릇과 군더더기가 섞여 있습니다.', '']
    for who, text in TRANSCRIPT:
        lines.append(f'[{t // 3600:02d}:{t % 3600 // 60:02d}:{t % 60:02d}] {who}: {text}')
        t += int(len(text) / 3.2) + rng.randint(3, 9)
    lines.append(f'[{t // 3600:02d}:{t % 3600 // 60:02d}:{t % 60:02d}] (녹음 끝)')
    (out / '회의녹취_0905.txt').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    set_mtime(out / '회의녹취_0905.txt', D.MEETING['date'], '16:40')
    return t


# ---------------------------------------------------------------------------
# 02 다운로드폴더_정리전
# 각 항목: 파일 이름, 날짜, 프로젝트, 문서 종류, 기대 새 이름, 내용 만드는 함수
# ---------------------------------------------------------------------------
def normalize_office(path):
    """docx·xlsx는 안이 zip이라 저장 시각이 들어간다. 시각을 고정해 다시 묶는다."""
    import re
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in items:
            if name == 'docProps/core.xml':
                data = re.sub(rb'(<dcterms:(created|modified)[^>]*>)[^<]*', rb'\g<1>2026-09-01T09:00:00Z', data)
            z.writestr(zipfile.ZipInfo(name, date_time=(2026, 9, 1, 9, 0, 0)), data, zipfile.ZIP_DEFLATED)


def make_docx(path, title, paras):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = '맑은 고딕'
    st.font.size = Pt(11)
    if title:
        doc.add_heading(title, level=1)
    for p in paras:
        doc.add_paragraph(p)
    doc.core_properties.created = FIXED
    doc.core_properties.modified = FIXED
    doc.core_properties.last_modified_by = ''
    doc.core_properties.author = ''
    doc.save(path)
    normalize_office(path)


def make_xlsx(path, sheet, rows, widths=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(list(r))
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, w in enumerate(widths or []):
        ws.column_dimensions[chr(65 + i)].width = w
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    wb.properties.creator = ''
    wb.save(path)
    normalize_office(path)


def make_image(path, title, lines, size=(1200, 900), bg=(236, 228, 214)):
    im = Image.new('RGB', size, bg)
    d = ImageDraw.Draw(im)
    W, H = size
    d.rectangle([W * 0.18, H * 0.2, W * 0.82, H * 0.8], fill=(200, 170, 130), outline=(120, 90, 60), width=6)
    d.rectangle([W * 0.36, H * 0.38, W * 0.64, H * 0.6], fill=(250, 248, 240), outline=(90, 70, 50), width=3)
    ft, fl = font(44, bold=True), font(30)
    tw = d.textlength(title, font=ft)
    d.text(((W - tw) / 2, H * 0.42), title, font=ft, fill=(40, 30, 20))
    y = H * 0.84
    for ln in lines:
        d.text((40, y), ln, font=fl, fill=(60, 50, 40))
        y += 40
    im.save(path, quality=86)


def make_screenshot(path, heading, items):
    """참고 사이트를 브라우저로 캡처한 것처럼 (가상의 가게 홈페이지)."""
    W, H = 1440, 900
    im = Image.new('RGB', (W, H), (255, 255, 255))
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, W, 48], fill=(232, 232, 232))
    d.rectangle([180, 10, W - 180, 38], fill=(255, 255, 255))
    d.text((200, 12), 'example-bakery-site.test', font=font(20), fill=(90, 90, 90))
    d.text((80, 110), heading, font=font(52, bold=True), fill=(30, 30, 30))
    x = 80
    for it in items:
        d.rectangle([x, 260, x + 400, 620], fill=(225, 214, 196))
        d.text((x + 20, 640), it, font=font(28, bold=True), fill=(40, 40, 40))
        x += 440
    d.text((80, 760), '메뉴    가게 소개    오시는 길    공지', font=font(28), fill=(80, 80, 80))
    im.save(path)


DOWNLOADS = []


def dl(name, date, project, doctype, newname, maker, note=''):
    DOWNLOADS.append(dict(name=name, date=date, project=project, doctype=doctype, newname=newname, maker=maker, note=note))


P_GIFT, P_MENU, P_WEB, P_ADMIN, P_OPS = '01_추석선물세트', '02_메뉴판리뉴얼', '03_홈페이지', '04_행정·세무', '05_운영·거래처'

GIFT_COPY_V1 = ['느린오븐의 추석 선물세트를 소개합니다.', '쿠키 세트 26,000원 / 파운드 세트 40,000원 / 종합 세트 52,000원',
                '천천히 구운 마음을 담았습니다.']
GIFT_COPY_FINAL = ['느린오븐 추석 선물세트', '쿠키 세트 28,000원: 버터·얼그레이·초코칩 쿠키 12개',
                   '파운드 세트 42,000원: 무화과 파운드케이크 1개와 쿠키 6개',
                   '예약 마감 9월 19일(토), 매장 수령 9월 22일(화)~24일(목)', '천천히 구운 마음을 담았습니다.']

dl('최종_진짜최종.docx', '2026-09-06', P_GIFT, '안내문구', '2026-09-06_추석선물세트_안내문구_v2.docx',
   lambda p: make_docx(p, '추석 선물세트 안내 문구', GIFT_COPY_FINAL), '최종본')
dl('최종.docx', '2026-08-29', P_GIFT, '안내문구', '2026-08-29_추석선물세트_안내문구_v1.docx',
   lambda p: make_docx(p, '추석 선물세트 안내 문구', GIFT_COPY_V1), '구버전 (세 종류, 작년 가격)')
dl('견적서(1).pdf', '2026-08-31', P_GIFT, '견적서', '2026-08-31_추석선물세트_상자견적_300개_v1.pdf',
   lambda p: pdf_doc(p, '견 적 서', [('p', f'수신: {D.SHOP} 귀하\n견적일: 2026-08-31'), ('p', '공급자: 상자나라   사업자번호 000-00-00301   전화 02-0000-3301'),
                                     ('table', [('품목', '수량', '금액(원)'), ('크라프트 선물상자 (중, 인쇄 없음)', '300', '540,000')], (100, 25, 40)),
                                     ('h', '합계: 540,000원 (VAT 별도)')], ['제작 기간은 주문 후 10일입니다.']), '구버전 (300개)')
dl('견적서(2).pdf', '2026-09-02', P_GIFT, '견적서', '2026-09-02_추석선물세트_상자견적_200개_v2.pdf',
   lambda p: pdf_doc(p, '견 적 서', [('p', f'수신: {D.SHOP} 귀하\n견적일: 2026-09-02 (수정)'), ('p', '공급자: 상자나라   사업자번호 000-00-00301   전화 02-0000-3301'),
                                     ('table', [('품목', '수량', '금액(원)'), ('크라프트 선물상자 (중, 인쇄 없음)', '200', '380,000')], (100, 25, 40)),
                                     ('h', '합계: 380,000원 (VAT 별도)')], ['제작 기간은 주문 후 10일입니다.']), '최종 (200개, 회의에서 이걸로 결정)')
dl('견적서.pdf', '2026-08-26', P_MENU, '견적서', '2026-08-26_메뉴판리뉴얼_인쇄견적.pdf',
   lambda p: pdf_doc(p, '견 적 서', [('p', f'수신: {D.SHOP} 귀하\n견적일: 2026-08-26'), ('p', '공급자: 또박또박 인쇄   사업자번호 000-00-00113   전화 02-0000-1113'),
                                     ('table', [('품목', '수량', '금액(원)'), ('메뉴판 A3 컬러 출력', '10', '35,000'), ('코팅', '10', '15,000')], (100, 25, 40)),
                                     ('h', '합계: 50,000원 (VAT 포함)')]))
dl('IMG_4821.jpg', '2026-09-01', P_GIFT, '시안사진', '2026-09-01_추석선물세트_상자시안_A.jpg',
   lambda p: make_image(p, '시안 A', ['손글씨 라벨, 크라프트 끈'], ))
dl('IMG_4822.jpg', '2026-09-01', P_GIFT, '시안사진', '2026-09-01_추석선물세트_상자시안_B.jpg',
   lambda p: make_image(p, '시안 B', ['단정한 라벨, 크라프트 끈'], bg=(226, 230, 222)))
dl('IMG_4830.jpg', '2026-09-03', P_WEB, '사진', '2026-09-03_홈페이지_매장외관.jpg',
   lambda p: make_image(p, '느린오븐', ['매장 외관 (홈페이지 가게 소개용)'], bg=(210, 220, 228)))
dl('스크린샷 2026-09-03 오후 3.12.45.png', '2026-09-03', P_WEB, '참고자료', '2026-09-03_홈페이지_참고사이트_1.png',
   lambda p: make_screenshot(p, '오늘 구운 빵', ['식빵', '깜파뉴', '스콘']))
dl('스크린샷 2026-09-03 오후 3.14.02.png', '2026-09-03', P_WEB, '참고자료', '2026-09-03_홈페이지_참고사이트_2.png',
   lambda p: make_screenshot(p, '예약 주문 안내', ['케이크', '선물세트', '단체 주문']))
dl('제목없음.txt', '2026-09-04', P_WEB, '메모', '2026-09-04_홈페이지_메뉴구성메모.txt',
   lambda p: p.write_text('홈페이지 메뉴 구성 (생각나는 대로)\n- 메인: 오늘 구운 빵\n- 가게 소개: 느리게 굽는 이유\n- 메뉴: 빵, 구움과자, 선물세트\n- 오시는 길\n- 공지: 휴무일\n예약 기능 꼭 넣기\n', encoding='utf-8'))
dl('새 문서 3.docx', '2026-08-21', P_ADMIN, '안내문', '2026-08-21_행정_위생교육일정메모.docx',
   lambda p: make_docx(p, None, ['위생교육 일정 메모', '올해 정기 위생교육: 온라인 수강, 9월 30일까지', '수료증 받으면 매장 서류철에 같이 보관']))
dl('새 문서 4.docx', '2026-08-28', P_MENU, '안내문구', '2026-08-28_메뉴판리뉴얼_신메뉴설명.docx',
   lambda p: make_docx(p, '신메뉴 설명', ['무화과 깜파뉴: 무화과를 듬뿍 넣고 하루 숙성한 반죽', '흑임자 스콘: 고소한 흑임자와 버터', '얼그레이 쿠키: 찻잎을 갈아 넣은 쿠키']))
dl('메뉴판_수정.xlsx', '2026-09-07', P_MENU, '가격표', '2026-09-07_메뉴판리뉴얼_가격표_v2.xlsx',
   lambda p: make_xlsx(p, '가격표', [('메뉴', '가격(원)', '비고'), ('우유식빵', 5500, ''), ('무화과 깜파뉴', 7500, '신메뉴'), ('흑임자 스콘', 3800, '신메뉴'),
                                    ('버터쿠키', 2200, '500원 인상'), ('얼그레이 쿠키', 2400, '신메뉴')], [18, 12, 14]), '최신')
dl('가격표.xlsx', '2026-08-12', P_MENU, '가격표', '2026-08-12_메뉴판리뉴얼_가격표_v1.xlsx',
   lambda p: make_xlsx(p, 'Sheet1', [('메뉴', '가격(원)'), ('우유식빵', 5500), ('버터쿠키', 1700), ('초코칩 쿠키', 1900)], [18, 12]), '구버전')
dl('다운로드.pdf', '2026-09-14', P_ADMIN, '수료증', '2026-09-14_행정_위생교육수료증.pdf',
   lambda p: pdf_doc(p, '수 료 증', [('p', f'성명: {D.OWNER}\n업소명: {D.SHOP}'), ('gap', 6),
                                     ('p', '위 사람은 2026년도 식품위생 정기교육(온라인)을\n이수하였기에 이 증서를 드립니다.'), ('gap', 10),
                                     ('p', '2026년 9월 14일\n가상 식품위생교육원 (이 문서는 실습용 가상 문서입니다)')]))
dl('document.pdf', '2026-09-10', P_ADMIN, '안내문', '2026-09-10_세무_부가세신고준비안내.pdf',
   lambda p: pdf_doc(p, '부가가치세 신고 준비 안내', [('p', f'{D.SHOP} 대표님께'), ('gap', 3),
                                                  ('p', '다음 신고를 위해 아래 자료를 미리 모아 주세요.\n- 카드 매출 내역\n- 사업용 카드·현금 영수증 (식비, 교통비 등 경비)\n- 재료 매입 거래명세서'),
                                                  ('gap', 3), ('p', '바른장부 세무회계 (가상)   전화 02-0000-4401')]))
dl('계약서_스캔.pdf', '2026-08-18', P_ADMIN, '안내문', '2026-08-18_행정_임대료조정안내.pdf',
   lambda p: pdf_doc(p, '임대료 조정 안내', [('p', f'임차인: {D.OWNER} ({D.SHOP})\n임대인: 가상빌딩 관리사무소'), ('gap', 3),
                                            ('p', '2027년 1월 갱신부터 월 임대료 조정 협의를 요청드립니다.\n협의 기한: 2026년 10월 31일')]))
dl('사진.zip', '2026-09-08', P_MENU, '사진묶음', '2026-09-08_메뉴판리뉴얼_신메뉴사진.zip', None, '안에 신메뉴 사진 3장')
dl('선물세트_라벨.png', '2026-09-04', P_GIFT, '라벨시안', '2026-09-04_추석선물세트_라벨시안_v1.png',
   lambda p: make_image(p, '느린오븐 추석', ['라벨 시안 1차 (오세린)'], size=(1000, 600), bg=(244, 238, 226)))
dl('label_final.png', '2026-09-12', P_GIFT, '라벨시안', '2026-09-12_추석선물세트_라벨최종.png',
   lambda p: make_image(p, '느린오븐 추석', ['라벨 최종 (B 기본 + 이름 손글씨)'], size=(1000, 600), bg=(236, 240, 232)), '최종')
dl('메모.txt', '2026-08-25', P_GIFT, '메모', '2026-08-25_추석선물세트_일정메모.txt',
   lambda p: p.write_text('추석 준비 메모\n- 세트 종류 줄이기 (작년 세 종류 너무 힘듦)\n- 상자 미리 주문\n- 라벨 세린 님께 부탁\n- 예약은 표로 받기\n', encoding='utf-8'))
dl('회의록 0905.docx', '2026-09-05', P_OPS, '회의록', '2026-09-05_운영_직원회의메모.docx',
   lambda p: make_docx(p, '9/5 회의 메모 (간단히)', ['선물세트 두 종류', '상자 200개', '라벨 12일까지', '홈페이지 15일까지 업체 결정']))
dl('거래처 연락처.xlsx', '2026-08-14', P_OPS, '연락처', '2026-08-14_운영_거래처연락처.xlsx',
   lambda p: make_xlsx(p, '거래처', [('거래처', '품목', '연락처'), ('고운가루 상회', '밀가루', '02-0000-5501'), ('상자나라', '포장 상자', '02-0000-3301'),
                                    ('또박또박 인쇄', '인쇄', '02-0000-1113'), ('포장마켓', '포장재 (온라인)', '-')], [18, 16, 16]))
dl('photo_2026-09-10.jpg', '2026-09-10', P_MENU, '사진', '2026-09-10_메뉴판리뉴얼_무화과깜파뉴.jpg',
   lambda p: make_image(p, '무화과 깜파뉴', ['신메뉴 사진 (메뉴판용)'], bg=(232, 222, 206)))
dl('invoice_0912.pdf', '2026-09-12', P_OPS, '거래명세서', '2026-09-12_운영_밀가루거래명세서.pdf',
   lambda p: pdf_doc(p, '거래명세서', [('p', f'공급받는 자: {D.SHOP}\n공급자: 고운가루 상회   사업자번호 000-00-00501'),
                                       ('table', [('품목', '수량', '금액(원)'), ('강력분 20kg', '4', '168,000'), ('박력분 20kg', '2', '76,000')], (100, 25, 40)),
                                       ('h', '합계: 244,000원 (VAT 포함)')]))
dl('새 텍스트 문서.txt', '2026-09-04', P_WEB, '메모', '2026-09-04_홈페이지_도메인후보메모.txt',
   lambda p: p.write_text('도메인 후보\nslowoven-bakery.test\nneurin-oven.test\n(둘 다 실습용 가상 주소)\n', encoding='utf-8'))
dl('제목 없는 스프레드시트.xlsx', '2026-09-09', P_GIFT, '예약목록', '2026-09-09_추석선물세트_예약주문표.xlsx',
   lambda p: make_xlsx(p, '예약', [('예약번호', '이름', '연락처', '세트', '개수', '수령일'),
                                  ('G-001', '김가상', '010-0000-0001', '쿠키 세트', 2, '2026-09-22'),
                                  ('G-002', '이예시', '010-0000-0002', '파운드 세트', 1, '2026-09-23'),
                                  ('G-003', '박샘플', '010-0000-0003', '쿠키 세트', 3, '2026-09-24')], [10, 10, 16, 12, 6, 12]),
   '실습용 가상 이름과 번호')
dl('이력서_아르바이트.pdf', '2026-08-30', P_OPS, '지원서', '2026-08-30_운영_아르바이트지원서.pdf',
   lambda p: pdf_doc(p, '아르바이트 지원서', [('p', '이름: 최가상\n연락처: 010-0000-0009\n희망 근무: 주말 오전'), ('gap', 3),
                                            ('p', '제과 매장 경험 1년. 포장과 계산 가능합니다.\n(실습용 가상 문서입니다)')]))
dl('견적서(3).pdf', '2026-09-03', P_GIFT, '견적서', '2026-09-03_추석선물세트_리본견적.pdf',
   lambda p: pdf_doc(p, '견 적 서', [('p', f'수신: {D.SHOP} 귀하\n견적일: 2026-09-03'), ('p', '공급자: 포장마켓   사업자번호 000-00-00105'),
                                     ('table', [('품목', '수량', '금액(원)'), ('빨강 공단 리본 (50m)', '2', '24,000')], (100, 25, 40)),
                                     ('h', '합계: 24,000원 (VAT 포함)')]), '리본 색 보류라 아직 결정 안 됨')
dl('메뉴판_시안.png', '2026-09-06', P_MENU, '시안', '2026-09-06_메뉴판리뉴얼_메뉴판시안.png',
   lambda p: make_image(p, '메뉴판 시안', ['가격표 v2 반영'], size=(900, 1200), bg=(240, 236, 228)))


def build_downloads():
    out = DIRS['d']
    for f in DOWNLOADS:
        p = out / f['name']
        if f['name'] == '사진.zip':
            with zipfile.ZipFile(p, 'w', zipfile.ZIP_DEFLATED) as z:
                for nm, title in [('무화과깜파뉴.jpg', '무화과 깜파뉴'), ('흑임자스콘.jpg', '흑임자 스콘'), ('얼그레이쿠키.jpg', '얼그레이 쿠키')]:
                    tmp = OUT / '_tmp.jpg'
                    make_image(tmp, title, ['신메뉴 사진'])
                    set_mtime(tmp, '2026-09-08')
                    z.write(tmp, nm)
                    tmp.unlink()
        else:
            f['maker'](p)
        set_mtime(p, f['date'], '15:00')


# ---------------------------------------------------------------------------
def write_notice():
    (OUT / '읽어주세요.txt').write_text(
        '실습용 샘플 파일입니다.\n\n'
        f'- 가상의 동네 빵집 "{D.SHOP}"(대표 {D.OWNER})의 자료라는 설정입니다.\n'
        '- 모든 사람 이름, 가게 이름, 전화번호, 주소, 사업자번호는 지어낸 것입니다.\n'
        '- 01_영수증: 8~9월 경비 영수증 사진 → 지출표 만들기 (미션 A)\n'
        '- 02_다운로드폴더_정리전: 뒤죽박죽 다운로드 폴더 → 정리하기 (미션 B)\n'
        '- 03_견적서비교: 홈페이지 제작 견적 세 곳 → 비교표 만들기 (미션 C)\n'
        '- 04_회의녹취: 회의 받아쓰기 → 회의록과 할 일 표 (미션 D)\n\n'
        '원본을 지키고 싶으면 폴더를 복사해서 복사본을 Cowork에 맡기세요.\n', encoding='utf-8')
    set_mtime(OUT / '읽어주세요.txt', '2026-09-30')


def make_zip():
    zp = OUT / 'cowork-samples.zip'
    if zp.exists():
        zp.unlink()
    with zipfile.ZipFile(zp, 'w', zipfile.ZIP_DEFLATED) as z:
        z.write(OUT / '읽어주세요.txt', 'cowork-samples/읽어주세요.txt')
        for d in DIRS.values():
            for f in sorted(d.iterdir()):
                z.write(f, f'cowork-samples/{d.name}/{f.name}')
    return zp


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    for d in DIRS.values():
        d.mkdir(parents=True)
    build_receipts()
    build_downloads()
    build_quotes()
    secs = build_meeting()
    write_notice()
    zp = make_zip()
    # 실습 페이지와 함께 배포한다 (저장소에는 samples/ 쪽만 올린다)
    shutil.copy(zp, ROOT / 'practice-page/cowork-samples.zip')
    print(f'영수증 {len(list(DIRS["r"].iterdir()))}장, 다운로드 {len(list(DIRS["d"].iterdir()))}개, '
          f'견적 {len(list(DIRS["q"].iterdir()))}개, 녹취 {secs // 60}분 {secs % 60}초, zip {zp.stat().st_size // 1024} KB')


if __name__ == '__main__':
    main()
