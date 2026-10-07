"""샘플 파일과 정답지가 서로 맞는지 확인한다.

실행: tools/.venv/bin/python tools/verify_samples.py
샘플 파일을 직접 열어서 읽은 값과 정답지(answer-keys.md)의 숫자를 대조한다.
"""
import pathlib
import re
import sys
import zipfile

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

import gen_samples as G
import sample_data as D

ROOT = pathlib.Path(__file__).resolve().parent.parent
S = ROOT / 'samples'
KEY = (ROOT / 'facilitator/answer-keys.md').read_text(encoding='utf-8')
fails = []


def check(ok, msg):
    print(('통과 ' if ok else '실패 ') + msg)
    if not ok:
        fails.append(msg)


def won(n):
    return f'{n:,}원'


def pdf_text(p):
    return ''.join(page.extract_text() or '' for page in PdfReader(p).pages)


# 01 영수증: 파일이 모두 열리고, 정답지 합계가 데이터에서 다시 계산한 값과 같다
rdir = S / '01_영수증'
files = sorted(rdir.iterdir())
check(len(files) == len(D.RECEIPTS), f'영수증 사진 {len(files)}장 = 데이터 {len(D.RECEIPTS)}장')
for f in files:
    Image.open(f).verify()
check(True, '영수증 사진 모두 열림')
uniq = [r for r in D.RECEIPTS if 'dup_of' not in r]
total = sum(D.receipt_total(r) for r in uniq)
check(f'합계 {won(total)}' in KEY, f'정답지 영수증 합계 {won(total)}')
for r in uniq:
    check(won(D.receipt_total(r)) in KEY and r['file'] in KEY, f"정답지에 {r['file']} {won(D.receipt_total(r))}")
dup = [r for r in D.RECEIPTS if 'dup_of' in r]
check(len(dup) == 1 and dup[0]['file'] in KEY, '중복 의심 1건이 정답지에 있음')

# 02 다운로드: 파일 수, 열리는지, 함정 파일의 실제 내용
ddir = S / '02_다운로드폴더_정리전'
names = sorted(p.name for p in ddir.iterdir())
check(names == sorted(f['name'] for f in G.DOWNLOADS), f'다운로드 폴더 {len(names)}개 = 정의 {len(G.DOWNLOADS)}개')
for p in ddir.iterdir():
    if p.suffix == '.docx':
        Document(p)
    elif p.suffix == '.xlsx':
        load_workbook(p)
    elif p.suffix == '.pdf':
        pdf_text(p)
    elif p.suffix in ('.png', '.jpg'):
        Image.open(p).verify()
    elif p.suffix == '.zip':
        check(len(zipfile.ZipFile(p).namelist()) == 3, '사진.zip 안에 3장')
check(True, '다운로드 폴더 파일 모두 열림')
check('리본' in pdf_text(ddir / '견적서(3).pdf'), '견적서(3).pdf는 리본 견적')
check('수 료 증' in pdf_text(ddir / '다운로드.pdf') or '수료' in pdf_text(ddir / '다운로드.pdf'), '다운로드.pdf는 수료증')
check('300' in pdf_text(ddir / '견적서(1).pdf') and '200' in pdf_text(ddir / '견적서(2).pdf'), '상자 견적 300개 판과 200개 판')
old = ' '.join(p.text for p in Document(ddir / '최종.docx').paragraphs)
new = ' '.join(p.text for p in Document(ddir / '최종_진짜최종.docx').paragraphs)
check('26,000' in old and '28,000' in new, '최종.docx가 구버전, 최종_진짜최종.docx가 최신')
ws = load_workbook(ddir / '메뉴판_수정.xlsx').active
check(any(r[0] == '버터쿠키' and r[1] == 2200 for r in ws.iter_rows(values_only=True)), '메뉴판_수정.xlsx 버터쿠키 2,200원')
for f in G.DOWNLOADS:
    check(f['newname'] in KEY, f"정답지에 {f['name']} → {f['newname']}")

# 03 견적서: PDF에 적힌 금액이 정답지 계산에 쓰인 금액과 같다
qdir = S / '03_견적서비교'
for q in D.QUOTES:
    text = pdf_text(qdir / q['file']).replace(' ', '')
    for name, amount, _ in q['year']:
        if '×' in name:  # 유지보수는 월 단가로 적혀 있다
            unit = int(re.search(r'× ([\d,]+)원', name).group(1).replace(',', ''))
            check(f'{unit:,}' in text, f"{q['vendor']} PDF에 유지보수 월 {unit:,}원")
        else:
            check(f'{amount:,}' in text, f"{q['vendor']} PDF에 {name} {amount:,}원")
    check(f'**{won(D.quote_year_total(q))}**' in KEY, f"정답지 {q['vendor']} 1년 비용 {won(D.quote_year_total(q))}")

# 04 회의녹취: 결정과 할 일의 근거가 녹취에 있다
t = (S / '04_회의녹취/회의녹취_0905.txt').read_text(encoding='utf-8')
for who, *_ in D.MEETING['people']:
    check(t.count(f'] {who}:') >= 5, f'녹취에 {who}의 발언 5회 이상')
for word in ['이만 팔천', '사만 이천', '십구일', '이십이일', '이십사일', '이백 개', '십이일', '십일일', '구일', '십오일', '보류']:
    check(word in t, f'녹취에 "{word}"')
last = re.findall(r'\[(\d\d):(\d\d):(\d\d)\]', t)[-1]
mins = int(last[0]) * 60 + int(last[1])
check(18 <= mins <= 22, f'녹취 길이 약 {mins}분 (18~22분)')

# zip: 한글 이름이 UTF-8 표시로 들어 있고 모든 파일이 있다
z = zipfile.ZipFile(S / 'cowork-samples.zip')
infos = z.infolist()
check(all(i.flag_bits & 0x800 for i in infos if not i.filename.isascii()), 'zip의 한글 파일 이름에 UTF-8 표시 (Windows에서 안 깨짐)')
count = sum(1 for d in G.DIRS.values() for _ in d.iterdir()) + 1
check(len(infos) == count, f'zip 안 파일 {len(infos)}개 = 샘플 {count}개')
check(z.testzip() is None, 'zip 무결성')

print()
print('모두 통과' if not fails else f'실패 {len(fails)}건')
sys.exit(1 if fails else 0)
